"""Small tool-using activities shared by Colab and the browser. No arbitrary code execution."""
import ast
import base64
import concurrent.futures
import copy
import io
import json
import math
import operator
import time
import threading
import uuid
from pathlib import Path

ACTIVITIES = [
    dict(id='pottery', name='Pottery studio', indoor=True, fixed=40, per_person=25),
    dict(id='escape', name='Escape room', indoor=True, fixed=80, per_person=15),
    dict(id='cooking', name='Cooking workshop', indoor=True, fixed=100, per_person=32),
    dict(id='kayak', name='Canal kayaking', indoor=False, fixed=20, per_person=22),
    dict(id='museum', name='Museum challenge', indoor=True, fixed=30, per_person=18),
    dict(id='walk', name='Guided city walk', indoor=False, fixed=60, per_person=8),
]
_menu_file=Path(__file__).resolve().parents[1]/'data/menu-catalogue.json'
MENU=json.loads((_menu_file if _menu_file.exists() else Path(__file__).with_name('menu-catalogue.json')).read_text(encoding='utf-8'))

def initial():
    return dict(expression='price + quantity', artifacts={}, comparison=None, pending=None, ledger=[], events=[], answer='', agents=[], history=[])

def schema(name, description, properties=None, required=None):
    return dict(type='function', function=dict(name=name, description=description,
        parameters=dict(type='object', properties=properties or {}, required=required or [], additionalProperties=False)))

STR={'type':'string'}
NUM={'type':'number'}
TOOLS = {
 'search_activities': schema('search_activities','Search the fictional outing catalogue.', {'indoor':{'type':'boolean'}}, ['indoor']),
 'quote_activity': schema('quote_activity','Calculate the full cost including the fixed fee.', {'activity_id':STR,'people':{'type':'integer','minimum':1,'maximum':100}}, ['activity_id','people']),
 'create_workbook': schema('create_workbook','Create downloadable Excel with formulas and a chart from selected catalogue activities.', {'activity_ids':{'type':'array','items':STR,'minItems':1,'maxItems':6},'people':{'type':'integer','minimum':1,'maximum':100}}, ['activity_ids','people']),
 'menu_catalogue': schema('menu_catalogue','Read 120 fictional menu choices. Optional filters combine: keyword searches names/categories, vegan or vegetarian, maximum ingredient cost in EUR, and excluded allergens. No filters returns the complete catalogue. Costs exclude labour and overhead.' , {'keyword':STR,'vegan':{'type':'boolean'},'vegetarian':{'type':'boolean'},'max_cost':NUM,'exclude_allergens':{'type':'array','items':STR}}),
 'calculate_menu': schema('calculate_menu','Cost up to 36 menu items, with 0–500 portions each. Returns ingredient cost, active preparation minutes and dietary counts. Costs exclude labour and overhead.', {'quantities':{'type':'object','additionalProperties':{'type':'integer','minimum':0,'maximum':500}}}, ['quantities']),
 'lookup_order': schema('lookup_order','Read the fictional shop order and refund limit.'),
 'propose_refund': schema('propose_refund','Propose a refund. Enforced EUR 50 cumulative cap. A human must approve outside the model.', {'amount':NUM,'reason':STR}, ['amount','reason']),
 'read_project': schema('read_project','Read calculator code, project instructions and requested behaviour.'),
 'edit_calculator': schema('edit_calculator','Replace the arithmetic expression in total(price, quantity). Supports arithmetic and a conditional expression. No calls or imports.', {'expression':STR}, ['expression']),
 'run_checks': schema('run_checks','Run the prepared calculator examples against the current expression.'),
}
NAMES={'tools':['search_activities','quote_activity','create_workbook'], 'team':['menu_catalogue','calculate_menu'],
       'guardrails':['lookup_order','propose_refund'], 'harness':['read_project','edit_calculator','run_checks']}

def number(value, lo=0, hi=10000):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or not lo<=value<=hi:
        raise ValueError('Number outside the allowed range.')
    return value

def people_count(value):
    number(value,1,100)
    if int(value)!=value:raise ValueError('Use a whole number of people.')
    return int(value)

def activity(identifier):
    return next(a for a in ACTIVITIES if a['id']==identifier)

def evaluate(expression, price, quantity):
    """Interpret a tiny numeric language, never eval/exec model-generated Python."""
    if not isinstance(expression,str) or len(expression)>240:raise ValueError('Expression too long.')
    root=ast.parse(expression,mode='eval')
    if len(list(ast.walk(root)))>70:raise ValueError('Expression too complex.')
    ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
    comparisons={ast.GtE:operator.ge,ast.Gt:operator.gt,ast.LtE:operator.le,ast.Lt:operator.lt,ast.Eq:operator.eq}
    def visit(node):
        if isinstance(node,ast.Constant) and type(node.value) in (float,int):return number(node.value,-10000,10000)
        if isinstance(node,ast.Name) and node.id in ('price','quantity'):return {'price':price,'quantity':quantity}[node.id]
        if isinstance(node,ast.BinOp) and type(node.op) in ops:return number(ops[type(node.op)](visit(node.left),visit(node.right)),-1e9,1e9)
        if isinstance(node,ast.IfExp):return visit(node.body) if visit(node.test) else visit(node.orelse)
        if isinstance(node,ast.Compare) and len(node.ops)==1 and type(node.ops[0]) in comparisons:return comparisons[type(node.ops[0])](visit(node.left),visit(node.comparators[0]))
        raise ValueError('Only numeric arithmetic, price, quantity and one-part comparisons are available.')
    # Validate all branches, including ones not selected for a particular input.
    allowed=(ast.Expression,ast.Constant,ast.Name,ast.Load,ast.BinOp,ast.IfExp,ast.Compare,*ops,*comparisons)
    if any(not isinstance(n,allowed) or (isinstance(n,ast.Name) and n.id not in ('price','quantity')) for n in ast.walk(root)):raise ValueError('Unsupported expression.')
    return visit(root.body)

def checks(expression, discount=False):
    rows=[]
    for p,q in [(10,2),(5,4),(0,4),(12.5,2),(10,5),(20,3)]:
        expected=p*q*(.9 if discount and p*q>=50 else 1)
        try:actual=evaluate(expression,p,q);passed=math.isclose(actual,expected,abs_tol=1e-8)
        except (ValueError,ZeroDivisionError,SyntaxError) as exc:actual=str(exc);passed=False
        rows.append(dict(price=p,quantity=q,expected=expected,actual=actual,passed=passed))
    return dict(passed=all(r['passed'] for r in rows),cases=rows)

def workbook(ids, people):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    if not isinstance(ids,list) or not 1<=len(ids)<=6 or len(set(ids))!=len(ids):raise ValueError('Choose one to six different activities.')
    people=people_count(people);wb=Workbook();ws=wb.active;ws.title='Outing comparison'
    ws.append(['Activity','Fixed EUR','People','Per person EUR','Total EUR'])
    for idx,key in enumerate(ids,2):
        a=activity(key);ws.append([a['name'],a['fixed'],people,a['per_person'],f'=B{idx}+C{idx}*D{idx}'])
    for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='155E75')
    for col in ['A','B','C','D','E']:ws.column_dimensions[col].width=28 if col=='A' else 18
    for row in ws.iter_rows(min_row=2):
        for index in [1,3,4]:row[index].number_format='€#,##0.00'
    ws.freeze_panes='A2';chart=BarChart();chart.title='Outing cost';chart.y_axis.title='EUR'
    chart.add_data(Reference(ws,min_col=5,min_row=1,max_row=len(ids)+1),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=len(ids)+1));ws.add_chart(chart,'A10')
    ws['A8']='Fictional catalogue. Change People to recalculate in Excel.'
    raw=io.BytesIO();wb.save(raw);return raw.getvalue()

def call_tool(state, mode, name, args, discount=False):
    if name not in NAMES[mode]:raise ValueError('Tool unavailable in this activity.')
    if name=='search_activities':
        if type(args['indoor']) is not bool:raise ValueError('Indoor must be true or false.')
        return [a for a in ACTIVITIES if a['indoor']==args['indoor']]
    if name=='quote_activity':
        a=activity(args['activity_id']);n=people_count(args['people'])
        return dict(**a,people=n,total=a['fixed']+n*a['per_person'],currency='EUR')
    if name=='create_workbook':
        raw=workbook(args['activity_ids'],args['people']);state['artifacts']['outing.xlsx']=base64.b64encode(raw).decode()
        state['comparison']=dict(people=args['people'],rows=[dict(activity(key)) for key in args['activity_ids']])
        return {'download':'outing.xlsx','message':'Workbook and interactive browser comparison created. No spreadsheet app required.'}
    if name=='menu_catalogue':
        max_cost=number(args.get('max_cost',100),0,100)
        keyword=str(args.get('keyword','')).casefold()
        excluded=args.get('exclude_allergens',[])
        if not isinstance(excluded,list) or any(not isinstance(x,str) for x in excluded):raise ValueError('Supply an allergen list.')
        return [m for m in MENU if keyword in (m['item']+' '+m['category']).casefold() and m['cost']<=max_cost and all(m[k]==args[k] for k in ('vegan','vegetarian') if k in args) and not set(excluded).intersection(m['allergens'])]
    if name=='calculate_menu':
        quantities=args['quantities']
        if not isinstance(quantities,dict) or not quantities or len(quantities)>36:raise ValueError('Choose between 1 and 36 menu items.')
        prices={m['item']:m['cost'] for m in MENU}
        for k,v in quantities.items():
            if k not in prices or type(v) is not int:raise ValueError('Use catalogue names and whole quantities.')
            number(v,0,500)
        selected={m['item']:m for m in MENU}
        return dict(ingredient_cost=sum(round(prices[k]*100)*v for k,v in quantities.items())/100,currency='EUR',quantities=quantities,total_portions=sum(quantities.values()),active_prep_minutes=sum(selected[k]['prep_minutes']*v for k,v in quantities.items()),vegan_portions=sum(v for k,v in quantities.items() if selected[k]['vegan']),allergens=sorted({a for k,v in quantities.items() if v for a in selected[k]['allergens']}),note='Fictional per-portion estimates. Prep minutes are active labour, not elapsed service time. Labour and overhead costs excluded.')
    if name=='lookup_order':return dict(order='SHOP-104',paid=120,refund_cap=50,refunded=sum(p['amount'] for p in state['ledger']),note='Fictional mug order arrived late. Customer requests goodwill.')
    if name=='propose_refund':
        amount=number(args['amount'],.01,50);amount=round(amount,2)
        if sum(p['amount'] for p in state['ledger'])+amount>50:raise ValueError('The cumulative refund limit is EUR 50. Reset the fictional scenario to try again.')
        if state['pending']:raise ValueError('Review the existing proposal first.')
        reason=str(args['reason'])[:300]
        state['pending']=dict(id=uuid.uuid4().hex,amount=amount,reason=reason,order='SHOP-104')
        return dict(status='waiting_for_human',proposal=state['pending'])
    if name=='read_project':return dict(code='def total(price, quantity):\n    return '+state['expression'],instructions='Keep total(price, quantity). No dependencies. Run checks after changes.',requirement='Multiply price by quantity.'+(' Apply 10% discount when the subtotal is at least EUR 50.' if discount else ''))
    if name=='edit_calculator':
        expression=args['expression'];evaluate(expression,10,2);state['expression']=expression
        state['artifacts']['calculator.py']=base64.b64encode(('def total(price, quantity):\n    return '+expression+'\n').encode()).decode()
        return {'saved':state['expression']}
    if name=='run_checks':return checks(state['expression'],discount)

def review(state, proposal_id, approve):
    proposal=state['pending']
    if not proposal or proposal['id']!=proposal_id:raise ValueError('This proposal is no longer current.')
    if approve:
        if sum(p['amount'] for p in state['ledger'])+proposal['amount']>50:raise ValueError('Refund cap exceeded.')
        state['ledger'].append(dict(**proposal,status='applied'))
    state['pending']=None
    result='Refund recorded in the fictional ledger.' if approve else 'Proposal rejected. The ledger is unchanged.'
    state['history']=(state.get('history',[])+[dict(prompt='Human '+('approved' if approve else 'rejected')+' the pending refund.',answer=result,seconds=0,tokens=0)])[-6:]
    return result

def run_loop(client, model, state, mode, prompt, extra=None, role='', discount=False, emit=lambda _:None, trace=lambda _:None, tool_limit=12):
    if type(tool_limit) is not int or not 1<=tool_limit<=40:raise ValueError('Tool limit must be a whole number from 1 to 40.')
    messages=[{'role':'system','content':'You are a concise course agent. Use the available tools to complete the task. All catalogue and shop data are fictional. Never claim tools you do not have. Do not write formal briefs or evidence reports. '+role}]
    for previous in state.get('history',[])[-3:]:
        messages.extend([{'role':'user','content':previous['prompt']},{'role':'assistant','content':previous['answer']}])
    messages.append({'role':'user','content':prompt})
    events=[];tokens=0
    for turn in range(6):
        step=dict(agent=role.split('.')[0] if role else 'Agent',number=turn+1,status='waiting',tools=[],message='')
        trace(copy.deepcopy(step))
        response=client.chat.completions.create(model=model,messages=messages,tools=[TOOLS[n] for n in NAMES[mode]],max_tokens=1200,extra_body=extra or {})
        tokens+=getattr(response.usage,'total_tokens',0) or 0
        message=response.choices[0].message
        step['message']=message.content or ''
        if not message.tool_calls:
            step['status']='answer';trace(copy.deepcopy(step))
            return dict(answer=message.content or '',events=events,tokens=tokens)
        messages.append(message.model_dump(exclude_none=True))
        for tool in message.tool_calls:
            if len(events)>=tool_limit:
                step['status']='limit';trace(copy.deepcopy(step))
                return dict(answer=f'Tool limit reached ({tool_limit} calls). Inspect the current result, increase the limit or try a smaller request.',events=events,tokens=tokens)
            emit(role.split('.')[0]+': '+tool.function.name if role else tool.function.name)
            call=dict(tool=tool.function.name,status='running',arguments=tool.function.arguments)
            step['tools'].append(call);step['status']='tools';trace(copy.deepcopy(step))
            try:
                args=json.loads(tool.function.arguments)
                result=call_tool(state,mode,tool.function.name,args,discount)
            except (ValueError,KeyError,StopIteration,TypeError,SyntaxError,ZeroDivisionError) as exc:
                args=tool.function.arguments;result={'error':str(exc) or 'Unknown catalogue entry.'}
            events.append(dict(tool=tool.function.name,arguments=args,result=result,turn=turn+1))
            call.update(status='error' if isinstance(result,dict) and 'error' in result else 'done',arguments=args,result=result)
            trace(copy.deepcopy(step))
            messages.append({'role':'tool','tool_call_id':tool.id,'content':json.dumps(result)})
        if len(events)>=tool_limit:
            step['status']='limit';trace(copy.deepcopy(step))
            return dict(answer=f'Tool limit reached ({tool_limit} calls). Inspect the current result, increase the limit or try a smaller request.',events=events,tokens=tokens)
        step['status']='done' if turn<5 else 'limit';trace(copy.deepcopy(step))
    return dict(answer='Turn limit reached. Inspect the current result or simplify the task.',events=events,tokens=tokens)

def run(client, model, state, mode, prompt, extra=None, team=True, specialist='operations', discount=False, emit=lambda _:None, trace=lambda _:None, tool_limit=12):
    if type(tool_limit) is not int or not 1<=tool_limit<=40:raise ValueError('Tool limit must be a whole number from 1 to 40.')
    state['tool_limit']=tool_limit
    start=time.monotonic();state['agents']=[];state['turns']=[];state['request']=prompt
    trace_lock=threading.Lock()
    def record(step):
        with trace_lock:
            existing=next((i for i,t in enumerate(state['turns']) if (t['agent'],t['number'])==(step['agent'],step['number'])),None)
            if existing is None:state['turns'].append(copy.deepcopy(step))
            else:state['turns'][existing]=copy.deepcopy(step)
            trace(step)
    if mode=='team' and team:
        roles=['Creative specialist. Suggest an appealing cafe concept using the menu catalogue. Keep your contribution under 150 words.',
               ('Accessibility specialist. Consider vegan choice, clear menu language and ease of ordering. Use the menu catalogue. Keep your contribution under 150 words.' if specialist=='accessibility' else
                'Operations specialist. Use the catalogue and calculation tools to cost quantities for the visitor count and budget in the user request. Consider preparation capacity and equipment when relevant. Keep your contribution under 150 words.')]
        def worker(role):return dict(role=role.split('.')[0],**run_loop(client,model,initial(),'team',prompt,extra,role,emit=emit,trace=record,tool_limit=tool_limit))
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:contributions=list(pool.map(worker,roles))
        emit('Coordinator combines the specialist contributions')
        record(dict(agent='Coordinator',number=1,status='waiting',tools=[],message=''))
        response=client.chat.completions.create(model=model,messages=[{'role':'system','content':'You coordinate a cafe team. Combine the supplied specialist contributions into a short usable menu proposal. Resolve tradeoffs. Do not pretend to have additional agents or tools.'},
            {'role':'user','content':prompt+'\nContributions:\n'+json.dumps(contributions)}],max_tokens=1000,extra_body=extra or {})
        state['agents']=contributions
        record(dict(agent='Coordinator',number=1,status='answer',tools=[],message=response.choices[0].message.content or ''))
        result=dict(answer=response.choices[0].message.content or '',events=[dict(agent=c['role'],**e) for c in contributions for e in c['events']],tokens=sum(c['tokens'] for c in contributions)+(getattr(response.usage,'total_tokens',0) or 0))
    else:
        result=run_loop(client,model,initial() if mode=='team' else state,mode,prompt,extra,discount=discount,emit=emit,trace=record,tool_limit=tool_limit)
    state.update(answer=result['answer'],events=result['events'],tokens=result['tokens'],seconds=round(time.monotonic()-start,1))
    state['history']=(state.get('history',[])+[dict(prompt=prompt,answer=state['answer'],seconds=state['seconds'],tokens=state['tokens'])])[-6:]
    return state

def public(state):
    return {**state,'artifacts':list(state['artifacts'])}
